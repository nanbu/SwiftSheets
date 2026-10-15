"""Independent judge for the synthetic specimen. Does not evaluate formulas or render drawings."""
import json
from contextlib import nullcontext
import os
from pathlib import Path
import posixpath
import subprocess
import tempfile
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "Tests/SwiftSheetsTests/Fixtures/interop/consumer-contract.xlsx"
NS = {
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "xda": "http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray",
    "r": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def verify_metadata(package, cells, expected_part="xl/metadata/arrays.xml", expected_flags=("1", "0"),
                    expected_origins=(("D1", "1"), ("H1", "0"))):
    # Resolve the part through its relationship, then follow the full cm -> bk -> rc -> futureMetadata chain.
    # Checking only the presence of a dynamicArrayProperties element would miss a dangling index or wrong GUID.
    relationships = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
    metadata_relationships = [rel for rel in relationships if rel.get("Type") ==
                              "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata"]
    assert len(metadata_relationships) == 1
    relationship = metadata_relationships[0]
    assert relationship.get("TargetMode", "Internal") == "Internal"
    target = relationship.get("Target")
    assert target
    part = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
    assert part == expected_part
    content_types = ET.fromstring(package.read("[Content_Types].xml"))
    assert any(node.get("PartName") == "/" + part and node.get("ContentType") ==
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"
               for node in content_types)

    metadata = ET.fromstring(package.read(part))
    assert metadata.tag == "{" + NS["s"] + "}metadata"
    types = metadata.find("s:metadataTypes", NS)
    assert types is not None
    type_records = types.findall("s:metadataType", NS)
    assert int(types.get("count")) == len(type_records) == 1
    assert type_records[0].get("name") == "XLDAPR" and type_records[0].get("cellMeta") == "1"
    future = metadata.findall('s:futureMetadata[@name="XLDAPR"]', NS)
    assert len(future) == 1
    future_records = future[0].findall("s:bk", NS)
    assert int(future[0].get("count")) == len(future_records) == len(expected_flags)
    flags = []
    for record in future_records:
        extensions = record.findall("s:extLst/s:ext", NS)
        assert len(extensions) == 1
        assert extensions[0].get("uri") == "{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}"
        properties = extensions[0].find("xda:dynamicArrayProperties", NS)
        assert properties is not None
        assert properties.get("fCollapsed") == "0"
        flags.append(properties.get("fDynamic"))
    assert flags == list(expected_flags)

    cell_metadata = metadata.find("s:cellMetadata", NS)
    assert cell_metadata is not None
    cell_records = cell_metadata.findall("s:bk", NS)
    assert int(cell_metadata.get("count")) == len(cell_records) == len(expected_origins)
    for address, expected_flag in expected_origins:
        cell = cells.find(f'.//s:c[@r="{address}"]', NS)
        assert cell is not None
        index = int(cell.get("cm"))
        assert 1 <= index <= len(cell_records)
        records = cell_records[index - 1].findall("s:rc", NS)
        assert len(records) == 1
        type_index, value_index = int(records[0].get("t")), int(records[0].get("v"))
        assert 1 <= type_index <= len(type_records)
        assert type_records[type_index - 1].get("name") == future[0].get("name")
        assert 0 <= value_index < len(future_records)
        assert flags[value_index] == expected_flag


def verify(path):
    # openpyxl explicitly does not model shape groups; package checks below are independent XML assertions.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="DrawingML support is incomplete.*")
        warnings.filterwarnings("ignore", message="Workbook contains no stylesheet.*")
        formulas = openpyxl.load_workbook(path, data_only=False)
        values = openpyxl.load_workbook(path, data_only=True)
    sheet = formulas["Interop"]
    assert sheet["D1"].value.text == "=_xlfn._xlws.FILTER(A1:A3,A1:A3>1)"
    assert sheet["D1"].value.ref == "D1:D2"
    assert sheet["H1"].value.ref == "H1:H2"
    assert sheet["F1"].value == "=SUM(_xlfn.ANCHORARRAY(D1))"
    assert values["Interop"]["D1"].value == 2
    assert values["Interop"]["D2"].value == 3
    assert values["Interop"]["F1"].value == 5
    assert sheet.row_dimensions[5].hidden
    with ZipFile(path) as package:
        drawing = ET.fromstring(package.read("xl/drawings/drawing1.xml"))
        assert len(drawing.findall("xdr:oneCellAnchor/xdr:grpSp/xdr:grpSp/xdr:sp", NS)) == 1
        assert drawing.find("xdr:oneCellAnchor/xdr:from/xdr:colOff", NS).text == "31750"
        assert drawing.find("xdr:absoluteAnchor/xdr:pos", NS).get("x") == "31750"
        assert drawing.find("xdr:twoCellAnchor", NS).get("editAs") == "oneCell"
        assert drawing.find("xdr:absoluteAnchor/xdr:sp/xdr:spPr/a:xfrm", NS).get("rot") == "1800000"
        body = drawing.find("xdr:absoluteAnchor/xdr:sp/xdr:txBody", NS)
        assert body.find("a:bodyPr", NS).get("lIns") == "31750"
        assert body.find("a:p/a:r/a:t", NS).text == "日本語"
        cells = ET.fromstring(package.read("xl/worksheets/sheet1.xml"))
        verify_metadata(package, cells)
        assert package.read("xl/custom/opaque.bin") == b"SwiftSheets synthetic opaque part\x00\x01\xff"
    formulas.close()
    values.close()


def main():
    verify(SOURCE)
    requested_directory = os.environ.get("SWIFTSHEETS_CONSUMER_INTEROP_DIR")
    if requested_directory:
        Path(requested_directory).mkdir(parents=True, exist_ok=True)
    destination = nullcontext(requested_directory) if requested_directory else tempfile.TemporaryDirectory(prefix="swiftsheets-consumer-")
    with destination as directory:
        environment = dict(os.environ, SWIFTSHEETS_CONSUMER_INTEROP_DIR=directory)
        subprocess.run(["swift", "test", "--filter", "ConsumerInteropTests/writesVerificationWorkbook"],
                       cwd=ROOT, env=environment, check=True)
        saved = Path(directory) / "consumer-contract.xlsx"
        verify(saved)
        created = Path(directory) / "new-dynamic.xlsx"
        formulas = openpyxl.load_workbook(created, data_only=False)
        values = openpyxl.load_workbook(created, data_only=True)
        assert formulas.active["B2"].value.text == "=_xlfn.SEQUENCE(A1)"
        assert formulas.active["B2"].value.ref == "B2:B4"
        assert [values.active[f"B{row}"].value for row in (2, 3, 4)] == [1, 2, 3]
        with ZipFile(created) as package:
            cells = ET.fromstring(package.read("xl/worksheets/sheet1.xml"))
            verify_metadata(package, cells, "xl/metadata.xml", ("1",), (("B2", "1"),))
        formulas.close()
        values.close()
        with ZipFile(SOURCE) as before, ZipFile(saved) as after:
            for name in ["xl/drawings/drawing1.xml", "xl/metadata/arrays.xml", "xl/custom/opaque.bin"]:
                assert before.read(name) == after.read(name), name
    print(json.dumps({"openpyxl": openpyxl.__version__, "codec_and_package": "passed",
                      "viewport": "unverified", "native_excel_recalculation": "unverified"}))


if __name__ == "__main__":
    main()
